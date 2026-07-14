"""CLI: analyze command and error handling."""

import pytest
from openpyxl import load_workbook

import hydrotrends as ht
from hydrotrends import cli
from hydrotrends.api import ReportColumn
from hydrotrends.core.config import InputSpec
from hydrotrends.core.constants import UNIT_PAIRS


def _daily_csv():
    return str(ht.datasets.sample_path("tarbela_daily"))


def test_unit_report_path_suffixes_flow_and_volume_units():
    spec = InputSpec(path="tarbela.csv", resolution="daily")
    rc = ReportColumn(
        "Inflow_Cusecs", "Cusecs", volume_column="Vol_MAF", volume_unit_label="MAF"
    )
    path = cli._unit_report_path("out/report.xlsx", spec, rc, single=True)
    assert path.name == "report_Cusecs_MAF.xlsx"
    assert path.parent.name == "out"


def test_unit_report_path_without_volume_pairing():
    spec = InputSpec(path="tarbela.csv", resolution="daily")
    rc = ReportColumn("Inflow_Cusecs", "Cusecs")
    path = cli._unit_report_path("report.xlsx", spec, rc, single=True)
    assert path.name == "report_Cusecs.xlsx"


def test_single_report_column_keeps_literal_output_path(tmp_path):
    """A single ReportColumn (e.g. one flow unit) writes to the literal -o
    path unchanged -- only 2+ columns trigger the per-unit-suffix naming."""
    out = tmp_path / "single.xlsx"
    config = ht.Config(
        inputs=[InputSpec(path=_daily_csv(), resolution="daily")],
        flow_units=[ht.FlowUnit.CUSECS],
    )
    columns = [
        ReportColumn(
            cli._FLOW_COLUMN[unit],
            unit.value,
            volume_column=cli._VOLUME_COLUMN[UNIT_PAIRS[unit]],
            volume_unit_label=UNIT_PAIRS[unit].value,
        )
        for unit in config.flow_units
    ]
    assert len(columns) == 1
    path = cli._output_path(str(out), config.inputs[0], single=True)
    assert path == out


def test_analyze_single_input_writes_two_workbooks(tmp_path):
    """Default run (both flow units) writes one workbook per unit pairing --
    Cusecs+MAF and Cumecs+BCM -- restoring the source script's two-workbook
    deliverable, instead of one combined file."""
    out = tmp_path / "r.xlsx"
    rc = cli.main(
        ["analyze", _daily_csv(), "-o", str(out), "--value-column", "inflow_cusec"]
    )
    assert rc == 0
    assert not out.exists()  # the literal path is a base name, not a real file
    cusecs_out = tmp_path / "r_Cusecs_MAF.xlsx"
    cumecs_out = tmp_path / "r_Cumecs_BCM.xlsx"
    assert cusecs_out.exists() and cumecs_out.exists()
    for path in (cusecs_out, cumecs_out):
        assert "Cover" in load_workbook(path).sheetnames


def test_analyze_writes_monthly_and_seasonal_volume_sheets(tmp_path):
    """Every flow unit is paired with its volume unit (UNIT_PAIRS) by default,
    so a plain run gets the monthly/hydro-season/met-season volume trend
    sheets too (both season schemes are on by default, per Config) -- split
    across the two per-unit workbooks."""
    out = tmp_path / "r.xlsx"
    cli.main(
        ["analyze", _daily_csv(), "-o", str(out), "--value-column", "inflow_cusec"]
    )
    maf_sheets = load_workbook(tmp_path / "r_Cusecs_MAF.xlsx").sheetnames
    bcm_sheets = load_workbook(tmp_path / "r_Cumecs_BCM.xlsx").sheetnames
    assert "Monthly_Trends_MAF" in maf_sheets
    assert "Hydro_Season_Trends_MAF" in maf_sheets
    assert "Met_Season_Trends_MAF" in maf_sheets
    assert "Monthly_Trends_BCM" in bcm_sheets
    assert "Hydro_Season_Trends_BCM" in bcm_sheets
    assert "Met_Season_Trends_BCM" in bcm_sheets


def test_10daily_compact(tmp_path):
    out = tmp_path / "t.xlsx"
    rc = cli.main(
        [
            "analyze",
            str(ht.datasets.sample_path("tarbela_10daily")),
            "-o",
            str(out),
            "-r",
            "10daily",
            "--value-column",
            "Inflow_1000Cusecs",
            "--value-scale",
            "1000",
            "--date-format",
            "dekad_compact",
        ]
    )
    assert rc == 0
    assert (tmp_path / "t_Cusecs_MAF.xlsx").exists()
    assert (tmp_path / "t_Cumecs_BCM.xlsx").exists()


def test_missing_file_exits_1(tmp_path):
    with pytest.raises(SystemExit) as exc:
        cli.main(["analyze", "/no/such.csv", "-o", str(tmp_path / "x.xlsx")])
    assert exc.value.code == 1


def test_no_input_exits_1(tmp_path):
    with pytest.raises(SystemExit) as exc:
        cli.main(["analyze", "-o", str(tmp_path / "x.xlsx")])
    assert exc.value.code == 1
