"""Analysis orchestration and end-to-end report generation."""

import math

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import hydrotrends as ht
from hydrotrends.api import analyze_by_period, analyze_series
from hydrotrends.core.config import SeasonScheme
from hydrotrends.core.constants import (
    COL_DATE,
    COL_DAY,
    COL_FLOW_CUSECS,
    COL_HYDRO_YEAR,
    COL_MONTH,
    COL_PERIOD,
    COL_VOL_MAF,
    COL_YEAR,
    HYDRO_MONTHS,
    TimeResolution,
)
from hydrotrends.viz.reports import STAT_GROUPS


def _synthetic_hydro_years(
    n_years=10, start_year=2010, base_flow=50_000.0, step=1000.0
):
    """``n_years`` complete hydrological years with a clear upward flow trend."""
    rng = np.random.default_rng(7)
    frames = []
    for i, year in enumerate(range(start_year, start_year + n_years)):
        dates = pd.date_range(f"{year}-04-01", f"{year + 1}-03-31")
        flow = base_flow + step * i + rng.normal(0, 50, len(dates))
        frames.append(pd.DataFrame({COL_DATE: dates, COL_FLOW_CUSECS: flow}))
    return ht.preprocess(pd.concat(frames, ignore_index=True), TimeResolution.DAILY)


def test_analyze_series_matches_individual_tests(trend_series):
    row = analyze_series(trend_series)
    assert row["trend"] == "increasing"
    assert math.isclose(row["p_value"], ht.mann_kendall(trend_series).p_value)
    assert math.isclose(row["sens_slope"], ht.sens_slope(trend_series).slope)


def test_all_row_keys_are_valid_schema_keys(step_series):
    row = analyze_series(step_series, years=range(2000, 2020))
    schema_keys = {c.key for g in STAT_GROUPS for c in g.columns}
    assert set(row) <= schema_keys


def test_change_point_years_populated(step_series):
    years = np.arange(2000, 2020)
    row = analyze_series(step_series, years=years)
    assert row["pettitt_cp_year"] == 2010
    assert row["cusum_cp_year"] == 2010 and row["bpcp_year"] == 2010


def test_n_below_4_descriptive_only():
    row = analyze_series([1.0, 2.0, 3.0])
    assert "trend" not in row and row["n"] == 3


def test_analyze_by_period_reindexes():
    recs = [
        {
            COL_PERIOD: p,
            COL_HYDRO_YEAR: 2000 + i,
            "flow": 100 + (3 * i if p == "Apr-01" else -2 * i),
        }
        for p in ("Apr-02", "Apr-01")
        for i in range(16)
    ]
    res = analyze_by_period(
        pd.DataFrame(recs), value_col="flow", order=["Apr-01", "Apr-02"]
    )
    assert list(res.index) == ["Apr-01", "Apr-02"]
    assert res.loc["Apr-01", "trend"] == "increasing"
    assert res.loc["Apr-02", "trend"] == "decreasing"


def test_analyze_preprocessed_ordering(daily_pre):
    res = ht.analyze_preprocessed(daily_pre, value_col=COL_FLOW_CUSECS)
    assert len(res) == 366 and res.index[0] == "Apr-01"


def test_analyze_preprocessed_change_point_uses_calendar_year_not_hydro_year():
    """Jan/Feb/Mar periods belong to HydroYear = Year - 1. The source script's
    daily/10-daily tables group the hydro-year-filtered frame by Period but
    always sort by and report calendar Year for change-point mapping (never
    HydroYear) -- confirmed straight from the source's Section 5 loop. This
    locks in that behaviour so it can't silently regress back to HydroYear."""
    frames = []
    for start_year in range(2010, 2020):
        dates = pd.date_range(f"{start_year}-04-01", f"{start_year + 1}-03-31")
        frames.append(
            pd.DataFrame({COL_DATE: dates, COL_FLOW_CUSECS: np.full(len(dates), 5.0)})
        )
    df = pd.concat(frames, ignore_index=True)

    # Step the Jan-15 value only: low for the first half of hydro years on
    # record, high for the second half.
    is_jan15 = (df[COL_DATE].dt.month == 1) & (df[COL_DATE].dt.day == 15)
    jan15_dates = df.loc[is_jan15, COL_DATE].sort_values()
    later_half = jan15_dates.iloc[len(jan15_dates) // 2 :]
    df.loc[df[COL_DATE].isin(later_half), COL_FLOW_CUSECS] = 9.0

    pre = ht.preprocess(df, TimeResolution.DAILY)
    result = ht.analyze_preprocessed(pre, value_col=COL_FLOW_CUSECS)
    row = result.loc["Jan-15"]

    jan15_rows = pre.hydro[
        (pre.hydro[COL_MONTH] == "Jan") & (pre.hydro[COL_DAY] == 15)
    ].sort_values(COL_YEAR)
    expected = ht.pettitt_test(jan15_rows[COL_FLOW_CUSECS].to_numpy())
    calendar_years = jan15_rows[COL_YEAR].to_numpy()
    hydro_years = jan15_rows[COL_HYDRO_YEAR].to_numpy()

    assert expected.index is not None
    assert row["pettitt_cp_year"] == calendar_years[expected.index]
    assert row["pettitt_cp_year"] != hydro_years[expected.index]  # guards the fix


def test_generate_report_end_to_end(daily_pre, tmp_path):
    out = ht.generate_report(
        daily_pre,
        tmp_path / "report.xlsx",
        columns=[ht.ReportColumn(COL_FLOW_CUSECS, "Cusecs")],
        title="Test Report",
    )
    assert out.exists()
    wb = load_workbook(out)
    assert "Cover" in wb.sheetnames and "Daily_Trends_Cusecs" in wb.sheetnames
    assert "10Daily_Trends_Cusecs" in wb.sheetnames
    ws = wb["Daily_Trends_Cusecs"]
    assert ws.cell(1, 1).value == "Daily Inflow Trend Analysis - Test Report [Cusecs]"
    assert ws.cell(2, 1).value == "Daily"
    assert ws.cell(4, 1).value == "Apr-01"  # hydro order


def test_generate_report_writes_daily_and_10daily_data_sheets(daily_pre, tmp_path):
    out = ht.generate_report(
        daily_pre,
        tmp_path / "report.xlsx",
        columns=[
            ht.ReportColumn(
                COL_FLOW_CUSECS,
                "Cusecs",
                volume_column=COL_VOL_MAF,
                volume_unit_label="MAF",
            )
        ],
        title="Test Report",
    )
    wb = load_workbook(out)
    for name in (
        "Daily_Data_Cal_Year_Cusecs",
        "Daily_Data_Hydro_Year_Cusecs",
        "Daily_Data_Met_Year_Cusecs",
        "Daily_Data_Cal_Year_MAF",
        "Daily_Data_Hydro_Year_MAF",
        "Daily_Data_Met_Year_MAF",
        # abbreviated Cal/Hydro/Met Year key: the full form would push these 3
        # past Excel's 31-char sheet-name limit (e.g. "..._Hydro_Year_Cusecs" = 35).
        "10Daily_Mean_Data_CY_Cusecs",
        "10Daily_Mean_Data_HY_Cusecs",
        "10Daily_Mean_Data_MY_Cusecs",
        "10Daily_Data_Cal_Year_MAF",
        "10Daily_Data_Hydro_Year_MAF",
        "10Daily_Data_Met_Year_MAF",
    ):
        assert name in wb.sheetnames, name
    assert all(len(name) <= 31 for name in wb.sheetnames)

    cal = wb["Daily_Data_Cal_Year_Cusecs"]
    assert cal.cell(2, 1).value == "Year"
    assert cal.cell(4, 1).value.isdigit()  # plain calendar year, no "-YY" suffix

    hydro = wb["Daily_Data_Hydro_Year_Cusecs"]
    assert hydro.cell(2, 1).value == "Hydro Year"
    assert "-" in hydro.cell(4, 1).value  # "YYYY-YY" label

    met = wb["Daily_Data_Met_Year_Cusecs"]
    assert met.cell(2, 1).value == "Met Year"
    assert "-" in met.cell(4, 1).value


def test_generate_report_10daily_input_has_no_daily_data_sheets(tendaily_pre, tmp_path):
    out = ht.generate_report(
        tendaily_pre,
        tmp_path / "report.xlsx",
        columns=[ht.ReportColumn(COL_FLOW_CUSECS, "Cusecs")],
        title="Test Report",
    )
    wb = load_workbook(out)
    assert not any(name.startswith("Daily_Data_") for name in wb.sheetnames)
    # abbreviated Cal/Hydro/Met Year key: "10Daily_Mean_Data_Hydro_Year_Cusecs"
    # would be 35 chars, over Excel's 31-char sheet-name limit.
    assert "10Daily_Mean_Data_HY_Cusecs" in wb.sheetnames
    assert all(len(name) <= 31 for name in wb.sheetnames)


# ── monthly / seasonal volume aggregation + trend analysis ──────────────────
def test_analyze_monthly_volumes_hydro_month_order_and_trend():
    pre = _synthetic_hydro_years()
    result = ht.analyze_monthly_volumes(pre.hydro, value_col=COL_VOL_MAF)
    assert list(result.index) == list(HYDRO_MONTHS)
    assert (result["n"] == 10).all()
    assert (result["trend"] == "increasing").all()
    assert (result["p_value"] < 0.05).all()


def test_analyze_hydro_seasonal_volumes_order_and_kharif_consistency():
    pre = _synthetic_hydro_years()
    result = ht.analyze_hydro_seasonal_volumes(pre.hydro, value_col=COL_VOL_MAF)
    assert [s.split()[0] for s in result.index] == [
        "Early",
        "Late",
        "Kharif",
        "Rabi",
        "Annual",
    ]
    assert (result["n"] == 10).all()
    assert (result["trend"] == "increasing").all()


def test_analyze_met_seasonal_volumes_order_and_trend():
    pre = _synthetic_hydro_years()
    result = ht.analyze_met_seasonal_volumes(pre.hydro, value_col=COL_VOL_MAF)
    assert [s.split()[0] for s in result.index] == [
        "Winter",
        "Spring",
        "Summer",
        "Monsoon",
        "Autumn",
        "Annual",
    ]
    # Interior years have n=10; Winter/Spring can lose an edge year to the
    # calendar-year-boundary/hydro-year-boundary effects documented on
    # met_seasonal_volumes, so only require "most years present, trend holds".
    assert (result["n"] >= 9).all()
    assert (result["trend"] == "increasing").all()


def test_analyze_monthly_volumes_below_min_is_descriptive_only():
    pre = _synthetic_hydro_years(n_years=3)
    result = ht.analyze_monthly_volumes(pre.hydro, value_col=COL_VOL_MAF)
    assert (result["n"] == 3).all()
    assert "trend" not in result.columns


def test_describe_monthly_and_seasonal_volumes_agree_with_analyze():
    pre = _synthetic_hydro_years()
    trend = ht.analyze_monthly_volumes(pre.hydro, value_col=COL_VOL_MAF)
    desc = ht.describe_monthly_volumes(pre.hydro, value_col=COL_VOL_MAF)
    assert list(desc.index) == list(HYDRO_MONTHS)
    pd.testing.assert_series_equal(desc["mean"], trend["mean"])

    seas_trend = ht.analyze_hydro_seasonal_volumes(pre.hydro, value_col=COL_VOL_MAF)
    seas_desc = ht.describe_hydro_seasonal_volumes(pre.hydro, value_col=COL_VOL_MAF)
    pd.testing.assert_series_equal(seas_desc["mean"], seas_trend["mean"])

    met_trend = ht.analyze_met_seasonal_volumes(pre.hydro, value_col=COL_VOL_MAF)
    met_desc = ht.describe_met_seasonal_volumes(pre.hydro, value_col=COL_VOL_MAF)
    pd.testing.assert_series_equal(met_desc["mean"], met_trend["mean"])


def test_generate_report_writes_volume_sheets_when_paired(tmp_path):
    pre = _synthetic_hydro_years()
    out = ht.generate_report(
        pre,
        tmp_path / "report.xlsx",
        columns=[
            ht.ReportColumn(
                COL_FLOW_CUSECS,
                "Cusecs",
                volume_column=COL_VOL_MAF,
                volume_unit_label="MAF",
            )
        ],
        title="Test Report",
    )
    wb = load_workbook(out)
    for name in (
        "Monthly_Trends_MAF",
        "Monthly Descriptive (MAF)",
        "Hydro_Season_Trends_MAF",
        "Hydro Season Descriptive (MAF)",
        "Met_Season_Trends_MAF",
        "Met Season Descriptive (MAF)",
    ):
        assert name in wb.sheetnames
    assert wb["Monthly_Trends_MAF"].cell(4, 1).value == "Apr"


def test_generate_report_omits_volume_sheets_without_pairing(daily_pre, tmp_path):
    out = ht.generate_report(
        daily_pre,
        tmp_path / "report.xlsx",
        columns=[ht.ReportColumn(COL_FLOW_CUSECS, "Cusecs")],
        title="Test Report",
    )
    wb = load_workbook(out)
    assert not any("Monthly_Trends" in s or "Season_Trends" in s for s in wb.sheetnames)


def test_generate_report_season_schemes_gating(tmp_path):
    pre = _synthetic_hydro_years()
    out = ht.generate_report(
        pre,
        tmp_path / "report.xlsx",
        columns=[
            ht.ReportColumn(
                COL_FLOW_CUSECS,
                "Cusecs",
                volume_column=COL_VOL_MAF,
                volume_unit_label="MAF",
            )
        ],
        title="Test Report",
        season_schemes=(SeasonScheme.CROPPING,),
    )
    sheets = load_workbook(out).sheetnames
    assert "Hydro_Season_Trends_MAF" in sheets
    assert "Met_Season_Trends_MAF" not in sheets
    assert "Monthly_Trends_MAF" in sheets  # not gated by season_schemes
