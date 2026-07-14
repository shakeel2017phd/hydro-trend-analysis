"""Tests for :mod:`hydrotrends.viz.reports` data-pivot sheet writers."""

from __future__ import annotations

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from hydrotrends.core.utils import hydro_year_label, met_year_label
from hydrotrends.stats.extended_descriptive import describe_extended_by
from hydrotrends.viz.reports import (
    EXTENDED_STATS_COLUMNS,
    write_annual_data_sheet,
    write_extended_stats_sheet,
    write_monthly_data_sheet,
    write_period_data_sheet,
    write_seasonal_data_sheet,
)

# The 15 data-sheet stats, in the order write_*_data_sheet lays them out
# (v26's HSTAT_LABELS order == _DESCRIPTIVE_COLUMNS with "N" dropped).
_STAT_HEADERS = (
    "Mean",
    "Median",
    "Std Dev",
    "CV (%)",
    "Min",
    "P10",
    "P25",
    "P75",
    "P90",
    "Max",
    "IQR",
    "Range",
    "Skewness",
    "Kurtosis",
    "Sum",
)


def test_hydro_year_label():
    assert hydro_year_label(2020) == "2020-21"
    assert hydro_year_label(1999) == "1999-00"


def test_met_year_label():
    assert met_year_label(2020) == "2019-20"
    assert met_year_label(2001) == "2000-01"


def _sheet() -> tuple[Workbook, Worksheet]:
    wb = Workbook()
    wb.remove(wb.active)
    return wb, wb.create_sheet("Test")


def test_write_period_data_sheet_layout_and_stats():
    wb, ws = _sheet()
    years = [2000, 2001, 2002]
    periods = ["Apr-01", "Apr-02", "May-01"]
    pivot = pd.DataFrame(
        {
            "Apr-01": [10.0, 20.0, 30.0],
            "Apr-02": [11.0, 21.0, 31.0],
            "May-01": [12.0, 22.0, 32.0],
        },
        index=years,
    )
    write_period_data_sheet(
        ws,
        pivot,
        title="Daily test [Cusecs]",
        period_order=periods,
        month_order=["Apr", "May"],
        row_label="Hydro Year",
        unit="Cusecs",
    )
    assert ws["A1"].value == "Daily test [Cusecs]"
    assert ws["A2"].value == "Hydro Year"
    # month merge header spans the 2 April periods; May starts right after
    assert ws["B2"].value == "April"
    assert ws["D2"].value == "May"
    # period row labels on row 3
    assert ws["B3"].value == "Apr-01"
    assert ws["D3"].value == "May-01"
    assert ws.cell(row=3, column=2 + len(periods)).value == "Mean"
    # data rows start at row 4, one per year
    assert ws["A4"].value == "2000"
    assert ws.cell(row=4, column=2).value == 10  # Cusecs -> rounded int
    # bottom column-statistics block starts right after the data rows
    bottom = 4 + len(years)
    assert ws.cell(row=bottom, column=1).value == "Mean"
    assert ws.cell(row=bottom + len(_STAT_HEADERS) - 1, column=1).value == "Sum"


def test_write_period_data_sheet_label_fn():
    wb, ws = _sheet()
    pivot = pd.DataFrame({"Apr-01": [10.0]}, index=[2000])
    write_period_data_sheet(
        ws,
        pivot,
        title="t",
        period_order=["Apr-01"],
        month_order=["Apr"],
        row_label="Hydro Year",
        unit="Cusecs",
        label_fn=hydro_year_label,
    )
    assert ws["A4"].value == "2000-01"


def test_write_monthly_data_sheet():
    wb, ws = _sheet()
    years = [2000, 2001]
    months = ["Apr", "May", "Jun"]
    pivot = pd.DataFrame(
        {"Apr": [1.0, 2.0], "May": [3.0, 4.0], "Jun": [5.0, 6.0]}, index=years
    )
    write_monthly_data_sheet(
        ws, pivot, title="Monthly test", month_order=months, label_fn=hydro_year_label
    )
    assert ws["A1"].value == "Monthly test"
    assert ws["A2"].value == "Hydro Year"
    assert ws["A3"].value == "(YYYY-YY)"
    assert ws["A4"].value == "2000-01"
    assert ws.cell(row=4, column=2).value == 1.0
    bottom = 4 + len(years)
    assert ws.cell(row=bottom, column=1).value == "Mean"


def test_write_seasonal_data_sheet_no_row_stats():
    wb, ws = _sheet()
    years = [2000, 2001, 2002]
    data = {
        "Kharif": pd.Series([5.0, 6.0, 7.0], index=years),
        "Rabi": pd.Series([1.0, 2.0, 3.0], index=years),
    }
    write_seasonal_data_sheet(
        ws,
        data,
        title="Seasonal test",
        col_headers=["Kharif", "Rabi"],
        sub_labels=["(Apr1-Sep30)", "(Oct1-Mar31)"],
        all_years=years,
        col_keys=["Kharif", "Rabi"],
        unit="MAF",
        label_fn=hydro_year_label,
    )
    assert ws["A2"].value == "Hydro Year"
    assert ws["A4"].value == "2000-01"
    assert ws.cell(row=4, column=2).value == 5.0
    bottom = 4 + len(years)
    assert ws.cell(row=bottom, column=1).value == "Mean"
    # column stats for Kharif mean of [5,6,7] == 6
    assert ws.cell(row=bottom, column=2).value == 6.0


def test_write_annual_data_sheet_highlight_cols():
    wb, ws = _sheet()
    years = [2000, 2001]
    ann_df = pd.DataFrame(
        {
            "Annual Vol\n(MAF)": [10.0, 20.0],
            "Anomaly\n(MAF)": [-5.0, 5.0],
            "Anomaly\n(%)": [-10.0, 10.0],
        },
        index=years,
    )
    write_annual_data_sheet(
        ws,
        ann_df,
        title="Annual test",
        data_cols=list(ann_df.columns),
        label_fn=hydro_year_label,
        highlight_cols=frozenset({"Anomaly\n(%)"}),
    )
    assert ws["A2"].value == "Hydro Year"
    assert ws["A3"].value == "(YYYY-YY)"
    assert ws["A4"].value == "2000-01"
    # Anomaly (MAF) is negative -> red-ish fill; Anomaly (%) is highlighted -> light red
    red_cell = ws.cell(row=4, column=3)
    pct_cell = ws.cell(row=4, column=4)
    assert red_cell.fill.fgColor.rgb.endswith("FFC7CE")
    assert pct_cell.fill.fgColor.rgb.endswith("FFE0E0")


def _flow_by_year_period():
    rng = np.random.default_rng(1)
    years = list(range(2000, 2010))
    periods = ["Apr-01", "Apr-02", "May-01"]
    rows = [
        {"Year": y, "Period": p, "Flow": rng.uniform(100, 500)}
        for y in years
        for p in periods
    ]
    return pd.DataFrame(rows), years, periods


def test_write_extended_stats_sheet_horizontal_by_year():
    df, years, periods = _flow_by_year_period()
    horiz = describe_extended_by(
        df, value_col="Flow", by="Year", expected_n=len(periods)
    )
    wb, ws = _sheet()
    write_extended_stats_sheet(
        ws, horiz, title="Horizontal test", index_label="Year", unit="Cusecs"
    )
    assert ws["A1"].value == "Horizontal test"
    assert ws["A2"].value == "Year  [Cusecs]"
    headers = [h for _f, h in EXTENDED_STATS_COLUMNS]
    assert ws.cell(row=2, column=2).value == headers[0] == "N"
    assert ws.cell(row=2, column=2 + len(headers) - 1).value == headers[-1]
    assert ws["A3"].value == "2000"
    n_col = 2  # "N" is the first stat column
    assert ws.cell(row=3, column=n_col).value == len(periods)
    completeness_col = 2 + [f for f, _h in EXTENDED_STATS_COLUMNS].index(
        "completeness_pct"
    )
    assert ws.cell(row=3, column=completeness_col).value == 100.0


def test_write_extended_stats_sheet_vertical_by_period_with_last_n():
    df, years, periods = _flow_by_year_period()
    vert = describe_extended_by(
        df,
        value_col="Flow",
        by="Period",
        sort_col="Year",
        expected_n=len(years),
        last_n_window=5,
    )
    wb, ws = _sheet()
    write_extended_stats_sheet(
        ws, vert, title="Vertical test", index_label="Period", unit="Cusecs"
    )
    assert ws["A3"].value in periods
    last_n_col = 2 + [f for f, _h in EXTENDED_STATS_COLUMNS].index("last_n_mean")
    # every period has a full 10-year record, so last-5-years mean is populated
    # (Cusecs renders as a whole number, matching every other flow-unit column)
    assert ws.cell(row=3, column=last_n_col).value != ""
